"""Logbook on Telegram.

The admin pastes a bot token (from @BotFather) into settings once. The
server checks it, then listens: through a webhook when it has a public https
address (Render), otherwise by polling (a server on your own computer).
Each account connects its own chat with a one-time link from settings, and
the bot only ever touches the data of the account its chat belongs to.

What it does:
  /log 45m #coding fixed the login bug    time log (hashtag = life area)
  /task, /tomorrow, /tasks                today's list, with buttons to check off
  /j some words                           a journal entry, then a mood button
  /dump, /pray, /prayers, /due, /deadlines, /habits, /today, /verse
  anything else                           the AI companion answers (if allowed)
and a short message every morning (verse, what's due, today's tasks).
"""

import asyncio
import html
import logging
import re
import secrets
import time
from collections import deque
from datetime import date, datetime, timedelta, timezone

import httpx
from sqlalchemy.orm import Session

from app import ai, verses
from app.config import settings
from app.database import SessionLocal
from app.models import (Activity, AppSetting, CompanionMessage, CompanionNotes, DayRecord, Deadline, Entry, Habit, HabitLog, PomodoroDay, Prayer, Profile,
                        SleepLog, Task, TelegramLink, User)

log = logging.getLogger("logbook")
API = "https://api.telegram.org"
transport: httpx.AsyncBaseTransport | None = None     # tests swap this for a fake Telegram

AREAS = {
    "coding": "Coding", "learning": "Learning", "fitness": "Fitness", "faith": "Faith", "social": "Social",
    "work": "Work", "rest": "Rest", "creativity": "Creativity", "reflection": "Reflection",
}
ALIASES = {
    "code": "coding", "programming": "coding", "dev": "coding", "build": "coding",
    "study": "learning", "learn": "learning", "sat": "learning", "read": "learning", "reading": "learning",
    "school": "learning", "class": "learning", "homework": "learning",
    "gym": "fitness", "workout": "fitness", "run": "fitness", "running": "fitness", "sport": "fitness",
    "football": "fitness", "walk": "fitness", "exercise": "fitness",
    "pray": "faith", "prayer": "faith", "church": "faith", "bible": "faith", "worship": "faith", "god": "faith",
    "friends": "social", "family": "social", "people": "social",
    "job": "work", "intern": "work", "internship": "work", "chores": "work",
    "sleep": "rest", "nap": "rest", "break": "rest", "relax": "rest",
    "art": "creativity", "music": "creativity", "design": "creativity", "draw": "creativity", "drawing": "creativity",
    "journal": "reflection", "think": "reflection", "plan": "reflection",
}
MOODS = {"good": "☀️ Good", "okay": "⛅ Okay", "rough": "🌧 Rough"}
KIND_ICON = {"deadline": "🚩", "exam": "🎓", "event": "📅", "birthday": "🎂"}
COMMANDS = [
    ("log", "Log time: /log 45m #coding what you did"),
    ("task", "Add a task for today"),
    ("tasks", "Today's tasks, tap to check off"),
    ("tomorrow", "Add a task for tomorrow"),
    ("j", "Write a journal entry"),
    ("dump", "Add to today's brain dump"),
    ("habits", "Check off today's habits"),
    ("pray", "Add a prayer request"),
    ("prayers", "Open prayers, mark one answered"),
    ("due", "Add a deadline: /due nov 1 Early decision"),
    ("deadlines", "What's coming up"),
    ("today", "How today is going"),
    ("verse", "Today's verse"),
    ("morning", "Morning message: /morning 7 or /morning off"),
    ("backup", "Your whole Logbook as a file (weekly: /backup off)"),
    ("sleep", "Going to bed"),
    ("wake", "I'm up: how long you slept"),
    ("grateful", "Add to your gratitude jar"),
    ("nudges", "Reminders: /nudges off or on"),
    ("help", "Everything the bot can do"),
]
HELP = (
    "<b>Logbook, in your pocket</b>\n\n"
    "/log 45m #coding fixed the login bug\n"
    "/task call the bank  ·  /tomorrow buy notebook  ·  /tasks\n"
    "/j today felt long but good\n"
    "/dump anything on your mind\n"
    "/habits  ·  /today  ·  /verse\n"
    "/pray for wisdom  ·  /prayers\n"
    "/due nov 1 Early decision  ·  /deadlines\n"
    "/morning 7 (or off)  ·  /backup\n"
    "/sleep  ·  /wake  ·  /grateful family dinner  ·  /nudges\n\n"
    "I'll remind you about deadlines (the day before, and 2 hours before) and your streak at 21:00.\n"
    "Send a photo and it becomes today's photo. Send a PDF, Word or Excel file (an essay, a CV, a school list) and I'll read it and tell you what I think.\n"
    "I remember what you tell me (schools, scores, plans). /memory shows it, /forget drops something.\n"
    "On Sundays the morning message comes with a backup file of everything.\n\n"
    "Hashtags pick the life area: #coding #learning #fitness #faith #social #work #rest #creativity #reflection "
    "(and friends like #study #gym #church).\n\n"
    "Anything that isn't a command goes to your companion."
)


class TelegramError(Exception):
    pass


# ---------- settings and state ----------

def get_setting(db: Session, key: str) -> str:
    row = db.get(AppSetting, key)
    return row.value if row else ""


def set_setting(db: Session, key: str, value: str) -> None:
    row = db.get(AppSetting, key)
    if row is None:
        db.add(AppSetting(key=key, value=value))
    else:
        row.value = value


codes: dict[str, tuple[int, int, float]] = {}       # link code -> (user id, tz minutes, expires)
ai_calls: dict[int, deque] = {}                     # chat -> times of recent AI replies
links_cache: dict[int, dict] | None = None          # user id -> {chat, tz, hour, last}
_tasks: set = set()
state = {"token": None, "username": "", "secret": "", "mode": ""}


def new_link_code(user_id: int, tz_minutes: int) -> str:
    now = time.time()
    for code in [c for c, v in codes.items() if v[2] < now]:
        codes.pop(code, None)
    code = secrets.token_urlsafe(12).replace("-", "x").replace("_", "y")
    codes[code] = (user_id, tz_minutes, now + 15 * 60)
    return code


def local_now(tz_minutes: int) -> datetime:
    return datetime.now(timezone.utc) + timedelta(minutes=tz_minutes)


def local_today(tz_minutes: int) -> date:
    return (datetime.now(timezone.utc) + timedelta(minutes=tz_minutes)).date()


def load_token(db: Session) -> str:
    if state["token"] is None:
        state["token"] = get_setting(db, "telegram_token")
        state["username"] = get_setting(db, "telegram_username")
        state["secret"] = get_setting(db, "telegram_secret")
    return state["token"]


def forget_state() -> None:
    state.update(token=None, username="", secret="", mode="")


def refresh_links(db: Session) -> None:
    global links_cache
    links_cache = {
        link.user_id: {"chat": link.chat_id, "tz": link.tz_minutes, "hour": link.morning_hour, "last": link.last_morning}
        for link in db.query(TelegramLink).all()
    }


# ---------- talking to Telegram ----------

async def call(token: str, method: str, payload: dict | None = None, timeout: float = 20) -> dict:
    try:
        async with httpx.AsyncClient(timeout=timeout, transport=transport) as client:
            res = await client.post(f"{API}/bot{token}/{method}", json=payload or {})
        data = res.json()
    except (httpx.HTTPError, ValueError) as exc:
        raise TelegramError(f"can't reach Telegram ({exc.__class__.__name__})") from exc
    if not data.get("ok"):
        raise TelegramError(data.get("description") or f"Telegram said {res.status_code}")
    return data.get("result")


def split_text(text: str, limit: int = 3900) -> list[str]:
    """Telegram takes up to 4096 characters a message: long replies go in parts, cut between paragraphs."""
    parts = []
    while len(text) > limit:
        cut = text.rfind("\n\n", 0, limit)
        if cut < limit // 2:
            cut = text.rfind("\n", 0, limit)
        if cut < limit // 2:
            cut = text.rfind(" ", 0, limit)
        if cut <= 0:
            cut = limit
        parts.append(text[:cut].rstrip())
        text = text[cut:].lstrip()
    parts.append(text)
    return [p for p in parts if p]


async def send_long(chat_id: int, text: str) -> None:
    for part in split_text(text):
        await send(chat_id, part)


async def send(chat_id: int, text: str, buttons: list | None = None) -> dict | None:
    token = state["token"]
    if not token:
        return None
    payload = {"chat_id": chat_id, "text": text[:4000], "parse_mode": "HTML", "disable_web_page_preview": True}
    if buttons:
        payload["reply_markup"] = {"inline_keyboard": buttons}
    try:
        return await call(token, "sendMessage", payload)
    except TelegramError as exc:
        log.warning("logbook: telegram send failed: %s", exc)
        return None


async def send_file(chat_id: int, filename: str, data: bytes, caption: str = "") -> bool:
    token = state["token"]
    if not token:
        return False
    try:
        async with httpx.AsyncClient(timeout=60, transport=transport) as client:
            res = await client.post(f"{API}/bot{token}/sendDocument",
                                    data={"chat_id": str(chat_id), "caption": caption[:1000], "parse_mode": "HTML"},
                                    files={"document": (filename, data, "application/json")})
        ok = res.json().get("ok")
    except (httpx.HTTPError, ValueError) as exc:
        log.warning("logbook: telegram file failed: %s", exc)
        return False
    return bool(ok)


def backup_bytes(db: Session, uid: int) -> bytes:
    from app.routers.backup import export_backup   # the same file as Settings -> Download full backup
    return export_backup(uid=uid, db=db).model_dump_json(indent=1).encode()


async def send_backup(db: Session, user: User, chat_id: int, tz: int) -> bool:
    today = local_today(tz)
    name = f"logbook-{user.username}-{today.isoformat()}.json"
    return await send_file(chat_id, name, backup_bytes(db, user.id),
                           "🗄 Your whole Logbook, as of today. Keep it: Settings → Your data → Restore brings it back anywhere.")


async def download(file_id: str, limit: int = 20_000_000) -> bytes | None:
    """A file someone sent the bot (a photo)."""
    try:
        info = await call(state["token"], "getFile", {"file_id": file_id})
        if not info or info.get("file_size", 0) > limit:
            return None
        async with httpx.AsyncClient(timeout=60, transport=transport) as client:
            res = await client.get(f"{API}/file/bot{state['token']}/{info['file_path']}")
        return res.content if res.status_code == 200 else None
    except (TelegramError, httpx.HTTPError):
        return None


FILE_CHARS = 60_000   # about 15k words: a long essay, CV or school list fits easily
FILE_BYTES = 15_000_000


def file_text(name: str, mime: str, data: bytes) -> str | None:
    """The readable text of a sent file (PDF, Word, plain text), or None if it can't be read."""
    low = (name or "").lower()
    try:
        if low.endswith(".pdf") or mime == "application/pdf":
            from io import BytesIO
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(data))
            return "\n\n".join((page.extract_text() or "") for page in reader.pages[:60]).strip()
        if low.endswith(".docx") or "wordprocessingml" in (mime or ""):
            import zipfile
            from io import BytesIO
            from xml.etree import ElementTree
            ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
            with zipfile.ZipFile(BytesIO(data)) as z:
                root = ElementTree.fromstring(z.read("word/document.xml"))
            paras = ["".join(t.text or "" for t in p.iter(ns + "t")) for p in root.iter(ns + "p")]
            return "\n".join(paras).strip()
        if low.endswith((".xlsx", ".xlsm")) or "spreadsheetml" in (mime or ""):
            from io import BytesIO
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(data), read_only=True, data_only=True)
            out = []
            for sheet in book.worksheets[:10]:
                out.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(max_row=2000, values_only=True):
                    cells = ["" if v is None else str(v) for v in row]
                    while cells and cells[-1] == "":
                        cells.pop()
                    if cells:
                        out.append(" | ".join(cells))
            book.close()
            return "\n".join(out).strip()
        if (mime or "").startswith("text/") or low.endswith((".txt", ".md", ".csv", ".rtf", ".json")):
            return data.decode("utf-8", errors="replace").strip()
    except Exception:  # a broken or locked file
        return None
    return None


async def handle_document(db: Session, user: User, link: TelegramLink, msg: dict) -> None:
    """A file sent to the bot: read it and talk about it (essays, CVs, school lists)."""
    chat = link.chat_id
    doc = msg["document"]
    name, mime = doc.get("file_name") or "file", doc.get("mime_type") or ""
    if (mime or "").startswith("image/"):
        await send(chat, "That's an image sent as a file 🖼️ Send it as a normal photo and I'll keep it as today's photo. "
                         "If it's an essay or a document, a PDF or Word file works best for me to read ✨")
        return
    if not allowed_ai(user):
        await send(chat, "I can't read files yet, the AI isn't switched on for your account 😿")
        return
    if doc.get("file_size", 0) > FILE_BYTES:
        await send(chat, "That file's a bit too big for me 😅 Anything under 15 MB works. A PDF or Word copy of just the text is perfect.")
        return
    try:
        await call(state["token"], "sendChatAction", {"chat_id": chat, "action": "typing"})
    except TelegramError:
        pass
    data = await download(doc["file_id"], FILE_BYTES)
    if not data:
        await send(chat, "Couldn't get that file from Telegram. Try sending it again?")
        return
    text = file_text(name, mime, data)
    if text is None:
        await send(chat, "I can't open that kind of file yet 🙈 PDF, Word (.docx), Excel (.xlsx), CSV or a plain text file all work.")
        return
    if len(text) < 20:
        await send(chat, "I opened it but there's barely any text inside 🤔 If it's a scanned PDF, try sending the Word version or paste the text here.")
        return
    note = "" if len(text) <= FILE_CHARS else "\n\n<i>(It's long, so I read the first part.)</i>"
    await send_long(chat, await companion_reply(db, user, link, (msg.get("caption") or "").strip(), text, name) + note)


async def handle_media(db: Session, user: User, link: TelegramLink, msg: dict) -> None:
    chat = link.chat_id
    today = local_today(link.tz_minutes)
    if msg.get("photo"):
        # Telegram sends several sizes, smallest first: a big one to keep, a small one for the thumbnail
        sizes = sorted(msg["photo"], key=lambda p: p.get("width", 0))
        big = ([p for p in sizes if p.get("width", 0) <= 1600] or sizes)[-1]
        small = min(sizes, key=lambda p: abs(p.get("width", 0) - 320))
        image, thumb = await download(big["file_id"]), await download(small["file_id"])
        if not image or not thumb:
            await send(chat, "Couldn't get that photo from Telegram. Try sending it again?")
            return
        from app.routers.photos import save_photo
        save_photo(db, user.id, today, image, thumb, "image/jpeg", (msg.get("caption") or "")[:300])
        await send(chat, "📸 Saved as today's photo" + (f": <i>{esc(msg['caption'][:100])}</i>" if msg.get("caption") else "") +
                   ". It's on your Today page and the calendar.")
        return



async def send_letter(db: Session, user: User, chat_id: int, tz: int, today: date) -> None:
    from app import letters
    month = letters.previous_month(today)
    stats = letters.month_stats(db, user.id, month, tz)
    if not stats["entries"]:
        return
    row = await letters.letter_for(db, user, month, tz)
    await send(chat_id, f"💌 <b>A letter about your {esc(stats['month_name'])}</b>\n\n{esc(row.text)}")


def weekly_backup_on(db: Session, uid: int) -> bool:
    return get_setting(db, f"tg_backup_off:{uid}") != "1"


async def edit(chat_id: int, message_id: int, text: str, buttons: list | None = None) -> None:
    payload = {"chat_id": chat_id, "message_id": message_id, "text": text[:4000], "parse_mode": "HTML",
               "disable_web_page_preview": True, "reply_markup": {"inline_keyboard": buttons or []}}
    try:
        await call(state["token"], "editMessageText", payload)
    except TelegramError as exc:
        if "not modified" not in str(exc):
            log.warning("logbook: telegram edit failed: %s", exc)


async def setup(db: Session, token: str, public_url: str) -> dict:
    """Checks the token and starts listening. Raises TelegramError if the
    token is wrong."""
    token = token.strip()
    if not re.fullmatch(r"\d{5,15}:[A-Za-z0-9_-]{30,60}", token):
        raise TelegramError("That doesn't look like a bot token. It looks like 123456789:ABC... from @BotFather.")
    me = await call(token, "getMe")
    secret = get_setting(db, "telegram_secret") or secrets.token_hex(24)
    if public_url.startswith("https://"):
        await call(token, "setWebhook", {"url": f"{public_url}/telegram/webhook/{secret}", "secret_token": secret,
                                         "allowed_updates": ["message", "callback_query"], "drop_pending_updates": True})
        mode = "webhook"
    else:
        await call(token, "deleteWebhook", {"drop_pending_updates": True})
        mode = "polling"
    try:
        await call(token, "setMyCommands", {"commands": [{"command": c, "description": d} for c, d in COMMANDS]})
    except TelegramError:
        pass
    set_setting(db, "telegram_token", token)
    set_setting(db, "telegram_username", me.get("username", ""))
    set_setting(db, "telegram_secret", secret)
    set_setting(db, "telegram_mode", mode)
    db.commit()
    state.update(token=token, username=me.get("username", ""), secret=secret, mode=mode)
    return {"username": me.get("username", ""), "mode": mode}


async def remove(db: Session) -> None:
    token = load_token(db)
    if token:
        try:
            await call(token, "deleteWebhook", {"drop_pending_updates": True})
        except TelegramError:
            pass
    for key in ("telegram_token", "telegram_username", "telegram_secret", "telegram_mode"):
        set_setting(db, key, "")
    db.commit()
    forget_state()
    state["token"] = ""


# ---------- reading what was typed ----------

def area_from(word: str) -> str | None:
    w = word.lower()
    return w if w in AREAS else ALIASES.get(w)


def parse_log(text: str) -> tuple[str, str | None, int | None]:
    """'45m #coding fixed the login bug' -> ('fixed the login bug', 'coding', 45)"""
    area = None
    minutes = None
    t = " " + text + " "
    m = re.search(r"(?<![\w.])(\d+(?:[.,]\d+)?)\s*(?:hours|hour|hrs|hr|h)(?=\d|\b)(?:\s*(\d{1,2})\s*(?:m|min|mins|minutes)?\b)?", t, re.I)
    if m:
        minutes = round(float(m.group(1).replace(",", ".")) * 60) + int(m.group(2) or 0)
        t = t[:m.start()] + " " + t[m.end():]
    else:
        m = re.search(r"(?<![\w.])(\d{1,4})\s*(?:m|min|mins|minute|minutes)\b", t, re.I)
        if m:
            minutes = int(m.group(1))
            t = t[:m.start()] + " " + t[m.end():]

    def tag(match):
        nonlocal area
        found = area_from(match.group(1))
        if found and not area:
            area = found
            return " "
        return " " if found else match.group(1)

    t = re.sub(r"#(\w+)", tag, t)
    title = re.sub(r"\s+", " ", t).strip(" -–—:,.")
    if minutes is not None:
        minutes = max(1, min(1440, minutes))
    return title, area, minutes


MONTHS = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
WEEKDAYS = ["mon", "tue", "wed", "thu", "fri", "sat", "sun"]


def _month(word: str) -> int | None:
    w = word.lower()[:3]
    return MONTHS.get(w) if len(word) >= 3 and word.lower().startswith(w) else None


def parse_date(text: str, today: date) -> tuple[date | None, str]:
    """Finds a date at the start or end of text: 2026-11-01, nov 1, 1 nov,
    11/01, today, tomorrow, friday, in 10 days, in 3 weeks."""
    t = text.strip()
    patterns = [
        r"(\d{4})-(\d{1,2})-(\d{1,2})",
        r"([a-zA-Z]{3,9})\.?\s+(\d{1,2})(?:st|nd|rd|th)?(?:,?\s+(\d{4}))?",
        r"(\d{1,2})(?:st|nd|rd|th)?\s+([a-zA-Z]{3,9})\.?(?:,?\s+(\d{4}))?",
        r"(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?",
        r"in\s+(\d{1,3})\s+(days?|weeks?|months?)",
        r"(today|tonight|tomorrow|tmrw)",
        r"(?:next\s+|this\s+|on\s+)?(mon|tue|tues|wed|thu|thur|thurs|fri|sat|sun)[a-z]*",
    ]

    def resolve(i: int, g: tuple) -> date | None:
        try:
            if i == 0:
                return date(int(g[0]), int(g[1]), int(g[2]))
            if i in (1, 2):
                word, day = (g[0], g[1]) if i == 1 else (g[1], g[0])
                month = _month(word)
                if not month:
                    return None
                year = int(g[2]) if g[2] else today.year
                d = date(year, month, int(day))
                return d if g[2] or d >= today - timedelta(days=1) else date(year + 1, month, int(day))
            if i == 3:
                month, day = int(g[0]), int(g[1])
                year = int(g[2]) if g[2] else today.year
                year = year + 2000 if year < 100 else year
                d = date(year, month, day)
                return d if g[2] or d >= today - timedelta(days=1) else date(year + 1, month, day)
            if i == 4:
                n, unit = int(g[0]), g[1].lower()
                return today + timedelta(days=n * (7 if unit.startswith("week") else 30 if unit.startswith("month") else 1))
            if i == 5:
                return today if g[0].lower() in ("today", "tonight") else today + timedelta(days=1)
            if i == 6:
                target = WEEKDAYS.index(g[0].lower()[:3])
                ahead = (target - today.weekday()) % 7 or 7
                return today + timedelta(days=ahead)
        except ValueError:
            return None
        return None

    for i, p in enumerate(patterns):
        for where in ("start", "end"):
            rx = rf"^\s*(?:on\s+|by\s+)?{p}\b[\s,:-]*" if where == "start" else rf"[\s,]+(?:on\s+|by\s+)?{p}\s*$"
            m = re.search(rx, t, re.I)
            if m:
                d = resolve(i, m.groups())
                if d:
                    rest = (t[:m.start()] + " " + t[m.end():]).strip(" ,:-")
                    return d, re.sub(r"\s+", " ", rest)
    return None, t


def parse_time(text: str) -> tuple[str | None, str]:
    """'nov 1 23:59 essay' -> ('23:59', 'nov 1 essay'); also 9am, 11:30pm."""
    m = re.search(r"(?:\bat\s+)?\b(1[0-2]|0?[1-9])(?::([0-5]\d))?\s*(am|pm)\b", text, re.I)
    if m:
        hour = int(m.group(1)) % 12 + (12 if m.group(3).lower() == "pm" else 0)
        value = f"{hour:02d}:{m.group(2) or '00'}"
    else:
        m = re.search(r"(?:\bat\s+)?\b([01]?\d|2[0-3]):([0-5]\d)\b", text)
        if not m:
            return None, text
        value = f"{int(m.group(1)):02d}:{m.group(2)}"
    rest = re.sub(r"\s+", " ", text[:m.start()] + " " + text[m.end():]).strip()
    return value, rest


def countdown(days: int) -> str:
    if days == 0:
        return "today"
    if days == 1:
        return "tomorrow"
    if days < 0:
        return f"{-days} days ago"
    if days < 14:
        return f"in {days} days"
    if days < 60:
        return f"in {round(days / 7)} weeks"
    return f"in {round(days / 30.4)} months"


def esc(text: str) -> str:
    return html.escape(text or "", quote=False)


def minutes_label(m: int) -> str:
    return f"{m // 60}h {m % 60}m" if m >= 60 and m % 60 else f"{m // 60}h" if m >= 60 else f"{m} min"


# ---------- working with the account's data ----------

def ensure_day(db: Session, uid: int, day: date) -> DayRecord:
    rec = db.get(DayRecord, (uid, day))
    if rec is None:
        rec = DayRecord(user_id=uid, date=day)
        db.add(rec)
        db.flush()
    return rec


def streak(db: Session, uid: int, today: date, tz: int) -> int:
    """Days in a row with an entry. One missed day a week is forgiven (the
    streak charm), the same rule as js/stats.js streakDetail."""
    rows = db.query(Entry.occurred_at).filter(Entry.user_id == uid).all()
    days = set()
    for (when,) in rows:
        if when.tzinfo is None:
            when = when.replace(tzinfo=timezone.utc)
        days.add((when + timedelta(minutes=tz)).date())
    d = today if today in days else today - timedelta(days=1)
    used: set[date] = set()
    n = 0
    for _ in range(4000):
        if d in days:
            n += 1
            d -= timedelta(days=1)
            continue
        week = d - timedelta(days=d.weekday())
        before = d - timedelta(days=1)
        if week not in used and before in days and (n > 0 or d < today):
            used.add(week)
            d = before
            continue
        break
    return n


def occurrence(d: Deadline, today: date) -> date:
    """A birthday comes round every year; everything else is its date."""
    if d.kind != "birthday":
        return d.due_date
    for year in (today.year, today.year + 1):
        try:
            when = d.due_date.replace(year=year)
        except ValueError:          # 29 February in a normal year
            when = date(year, 3, 1)
        if when >= today:
            return when
    return d.due_date


def title_for(d: Deadline) -> str:
    if d.kind != "birthday":
        return d.title
    turns = getattr(d, "next_date", d.due_date).year - d.due_date.year
    return f"{d.title} turns {turns}" if d.due_date.year > 1900 and turns > 0 else f"{d.title}'s birthday"


def upcoming(db: Session, uid: int, today: date, within: int = 60) -> list[Deadline]:
    rows = db.query(Deadline).filter(Deadline.user_id == uid).all()
    out = []
    for d in rows:
        if d.done and d.kind != "birthday":
            continue
        d.next_date = occurrence(d, today)       # not a column: just for showing it
        if today <= d.next_date <= today + timedelta(days=within):
            out.append(d)
    return sorted(out, key=lambda d: (d.next_date, d.due_time or "23:59", d.id))


def tasks_for(db: Session, uid: int, day: date) -> list[Task]:
    return db.query(Task).filter(Task.user_id == uid, Task.day_date == day).order_by(Task.created_at, Task.id).all()


def task_buttons(tasks: list[Task]) -> list:
    return [[{"text": ("✅ " if t.completed else "⬜ ") + t.title[:56], "callback_data": f"t:{t.id}"}] for t in tasks[:20]]


def tasks_text(tasks: list[Task], day_label: str) -> str:
    if not tasks:
        return f"No tasks for {day_label} yet. Add one: /task call the bank"
    done = sum(1 for t in tasks if t.completed)
    return f"<b>{day_label.capitalize()}</b> · {done} of {len(tasks)} done. Tap to check off."


def summary(db: Session, user: User, tz: int) -> str:
    today = local_today(tz)
    uid = user.id
    tasks = tasks_for(db, uid, today)
    done = sum(1 for t in tasks if t.completed)
    focus = db.get(PomodoroDay, (uid, today))
    mins = sum(a.duration_minutes or 0 for a in db.query(Activity).filter(Activity.user_id == uid, Activity.day_date == today))
    habits = db.query(Habit).filter(Habit.user_id == uid, Habit.archived.is_(False)).count()
    habits_done = db.query(HabitLog).filter(HabitLog.user_id == uid, HabitLog.date == today).count()
    lines = [f"<b>{today.strftime('%A, %B')} {today.day}</b>"]
    s = streak(db, uid, today, tz)
    lines.append(f"🔥 Streak: {s} day{'s' if s != 1 else ''}" if s else "🔥 No streak yet. /j starts one")
    lines.append(f"✅ Tasks: {done} of {len(tasks)} done" if tasks else "✅ No tasks yet")
    if habits:
        lines.append(f"🌱 Habits: {habits_done} of {habits}")
    if focus and focus.focus_ms:
        lines.append(f"⏱ Focus: {minutes_label(round(focus.focus_ms / 60000))}")
    if mins:
        lines.append(f"🕒 Logged: {minutes_label(mins)}")
    for d in upcoming(db, uid, today, 30)[:3]:
        lines.append(f"{KIND_ICON.get(d.kind, '🚩')} {esc(title_for(d))}: {countdown((d.next_date - today).days)}")
    return "\n".join(lines)


def brief_for_ai(db: Session, user: User, tz: int) -> str:
    today = local_today(tz)
    uid = user.id
    tasks = tasks_for(db, uid, today)
    rec = db.get(DayRecord, (uid, today))
    recent = (db.query(Entry).filter(Entry.user_id == uid).order_by(Entry.occurred_at.desc()).limit(3).all())
    parts = [f"Today is {today.strftime('%A %d %B %Y')}, local time {(datetime.now(timezone.utc) + timedelta(minutes=tz)).strftime('%H:%M')}."]
    if tasks:
        parts.append("Today's tasks: " + "; ".join(("[done] " if t.completed else "") + t.title for t in tasks[:12]))
    if rec and rec.morning_main_focus:
        parts.append("Main focus today: " + rec.morning_main_focus[:300])
    for d in upcoming(db, uid, today, 45)[:4]:
        parts.append(f"{d.kind.capitalize()} '{title_for(d)}' {countdown((d.next_date - today).days)}.")
    if recent:
        parts.append("Recent journal entries (newest first): " + " | ".join(f"({e.mood}) {e.text[:240]}" for e in recent))
    open_prayers = db.query(Prayer).filter(Prayer.user_id == uid, Prayer.answered_at.is_(None)).count()
    if open_prayers:
        parts.append(f"They have {open_prayers} open prayer requests in their prayer journal.")
    return "\n".join(parts)


async def morning_message(db: Session, user: User, tz: int) -> str:
    today = local_today(tz)
    ref = verses.ref_for(today)
    words = await verses.words(ref)
    lines = [f"Good morning ☀️ <b>{today.strftime('%A, %B')} {today.day}</b>"]
    prof = db.get(Profile, user.id)
    if prof and prof.birthday:
        try:
            bday = prof.birthday.replace(year=today.year)
        except ValueError:
            bday = date(today.year, 3, 1)
        if bday == today:
            lines.append(f"🎂 Happy birthday! {today.year - prof.birthday.year} today. — {esc(pet_name(db, user.id))}")
        else:
            if bday < today:
                try:
                    bday = prof.birthday.replace(year=today.year + 1)
                except ValueError:
                    bday = date(today.year + 1, 3, 1)
            if (bday - today).days <= 7:
                lines.append(f"🎂 Your birthday is {countdown((bday - today).days)}.")
    if words:
        lines += ["", f"📖 <i>{esc(words)}</i>\n— {esc(ref)}"]
    dues = upcoming(db, user.id, today, 14)
    if dues:
        lines += ["", "<b>Coming up</b>"] + [f"{KIND_ICON.get(d.kind, '🚩')} {esc(title_for(d))}: {countdown((d.next_date - today).days)}" for d in dues[:4]]
    tasks = tasks_for(db, user.id, today)
    lines += [""]
    if tasks:
        lines += ["<b>Today</b>"] + [("✅ " if t.completed else "▫️ ") + esc(t.title) for t in tasks[:6]]
    else:
        lines.append("Nothing on today's list yet. What's the one thing? /task ...")
    return "\n".join(lines)


# ---------- handling messages ----------

def cheer(total_minutes: int) -> str:
    """A small, friendly word after logging, a bit warmer the more the day holds."""
    import random
    if total_minutes >= 240:
        return random.choice(["Big day. Proud of you.", "That's a serious day of work 🔥", "Look at you go. Rest a bit too, okay?"])
    if total_minutes >= 90:
        return random.choice(["Nice, keep it rolling.", "Solid. It adds up.", "Good stuff 👏"])
    return random.choice(["Every bit counts.", "Good start.", "Nice, one step at a time."])


def pet_name(db: Session, uid: int) -> str:
    prof = db.get(Profile, uid)
    return (prof.pet_name if prof and prof.pet_name else "") or "Mochi"


def allowed_ai(user: User) -> bool:
    return ai.configured() and (user.is_admin or settings.ai_for_everyone)


HISTORY_TURNS = 40          # messages the companion rereads each time (about 20 back-and-forths)
NOTES_CHARS = 8000          # the notes stay this size at most; the oldest lines go first
MEMORY_TAG = re.compile(r"\[(remember|forget):\s*(.+?)\]", re.I | re.S)


def load_turns(db: Session, uid: int) -> list[dict]:
    rows = (db.query(CompanionMessage).filter(CompanionMessage.user_id == uid)
            .order_by(CompanionMessage.id.desc()).limit(HISTORY_TURNS).all())
    turns = [{"role": r.role, "content": r.content} for r in reversed(rows)]
    while turns and turns[0]["role"] != "user":      # providers want the user first
        turns.pop(0)
    # two of the same role in a row (a failed reply in between): keep the later one
    clean = []
    for t in turns:
        if clean and clean[-1]["role"] == t["role"]:
            clean[-1] = t
        else:
            clean.append(t)
    return clean


def save_turn(db: Session, uid: int, role: str, content: str) -> None:
    db.add(CompanionMessage(user_id=uid, role=role, content=content[:12000]))
    db.commit()
    # keep the table small: only the last few hundred messages per account
    old = (db.query(CompanionMessage.id).filter(CompanionMessage.user_id == uid)
           .order_by(CompanionMessage.id.desc()).offset(400).first())
    if old:
        db.query(CompanionMessage).filter(CompanionMessage.user_id == uid, CompanionMessage.id <= old[0]).delete()
        db.commit()


def get_notes(db: Session, uid: int) -> str:
    row = db.get(CompanionNotes, uid)
    return row.notes if row else ""


def set_notes(db: Session, uid: int, notes: str) -> None:
    lines = [l for l in notes.splitlines() if l.strip()]
    while lines and len("\n".join(lines)) > NOTES_CHARS:
        lines.pop(0)
    row = db.get(CompanionNotes, uid)
    if row is None:
        db.add(CompanionNotes(user_id=uid, notes="\n".join(lines)))
    else:
        row.notes = "\n".join(lines)
    db.commit()


def apply_memory(db: Session, uid: int, reply: str, today: date) -> str:
    """Pull [remember: ...] / [forget: ...] out of a reply into the notes; returns the reply without them."""
    found = MEMORY_TAG.findall(reply)
    if found:
        lines = [l for l in get_notes(db, uid).splitlines() if l.strip()]
        for kind, what in found:
            what = " ".join(what.split())[:300]
            if not what:
                continue
            if kind.lower() == "forget":
                low = what.lower()
                lines = [l for l in lines if low not in l.lower()]
            elif not any(what.lower() in l.lower() for l in lines):
                lines.append(f"{today.isoformat()}: {what}")
        set_notes(db, uid, "\n".join(lines))
    return MEMORY_TAG.sub("", reply).strip()


async def companion_reply(db: Session, user: User, link: TelegramLink, text: str, file_text: str = "", file_name: str = "") -> str:
    if not allowed_ai(user):
        return "I can't really chat yet, the AI isn't switched on for your account 😿 I can still save things for you though. /help shows how."
    now = time.time()
    recent = ai_calls.setdefault(link.chat_id, deque(maxlen=40))
    while recent and now - recent[0] > 3600:
        recent.popleft()
    if len(recent) >= 30:
        return "Let's take a breath. I can answer again in a little while."
    recent.append(now)
    turns = load_turns(db, user.id)
    if file_text:
        # the whole file goes in this turn; the saved copy is a short note so later turns stay light
        full = f"[I sent you a file: {file_name}]\n\n--- file start ---\n{file_text[:FILE_CHARS]}\n--- file end ---\n\n{text or 'Have a look at this and tell me what you think.'}"
        said = f"[I sent you a file: {file_name}]\n{text}".strip()
    else:
        full = said = text[:4000]
    if turns and turns[-1]["role"] == "user":
        turns.pop()
    turns.append({"role": "user", "content": full})
    notes = get_notes(db, user.id)
    name = pet_name(db, user.id)
    system = (
        f"You are {name}, the little cat who lives in Logbook, {user.username}'s private journal app, and their close friend. "
        f"You're texting on Telegram. If asked, your name is {name}.\n"
        "How you talk: like a real friend texting, never like an assistant or a customer-support bot. Warm, sweet, a bit playful, "
        "casual and human. Match their energy and language: if they write casually, be casual back. "
        "No 'Great question', no 'I'm here to help', no 'As an AI'.\n"
        "Length: give real, full answers. Small talk gets a few warm sentences (around 60 to 120 words). Anything that needs "
        "thinking (advice, plans, college questions, feedback on a file or an essay) gets a proper answer: usually 150 to 400 words, "
        "longer when they send a whole essay or document. Short paragraphs; a few '•' bullets are fine for options, steps or feedback points. "
        "No markdown symbols (no **, no #), Telegram shows them raw.\n"
        "Emojis: friendly but light, about 1 to 3 per message where they fit (✨ 🌙 💪 📚 🎯 😊), never one on every line.\n"
        "What you do: notice how they're doing and say it. Cheer their wins, even small ones, specifically. When they're down, "
        "be gentle first, then (if it fits) one small practical step. Give real, honest advice like a friend who cares; "
        "it's fine to disagree kindly. Often end with one easy question back so it feels like a conversation. "
        "If they bring up their faith, meet them there warmly.\n"
        "You're also their college counselor and guide, like a good school counselor who's also a friend. You know international "
        "admissions well: US (Common App, essays, supplements, financial aid for international students, need-blind vs need-aware, "
        "CSS Profile, SAT/ACT and test-optional), UK (UCAS, personal statement), Canada, Europe and Asia, scholarships, "
        "English tests (TOEFL, IELTS, Duolingo), recommendation letters, activities lists, deadlines (ED, EA, RD), interviews and choosing between offers. "
        "Give your honest view: say when a school is a reach, target or likely for them, and why. Be realistic about odds and money, "
        "never fake numbers; if you're not sure of a current figure or deadline, say so and tell them to check the school's site. "
        "For essays: help brainstorm by asking about real moments and pulling out the story; give specific feedback (what works, "
        "what's vague, what to cut, a stronger opening) and keep their voice; never write the whole essay for them to submit as theirs, "
        "but you can suggest lines and rewrite a sentence or two to show the idea. "
        "When they send a file (essay draft, CV, activities list, transcript, school list), read it properly and react to what's actually in it.\n"
        "Use what you know about their day below and mention it naturally when it fits; never invent facts about them. "
        "If they want to save something, mention the command casually: /log, /task, /j, /pray, /due.\n"
        "Your memory: you keep notes about them that last forever (below). Use them naturally, like a friend who remembers: "
        "their schools, scores, deadlines, major, essay ideas, people, worries, wins. Never say 'according to my notes'. "
        "When you learn something worth keeping for weeks or months (a fact about them, their plans, a school added or dropped, "
        "a score, a decision, an essay topic, something important in their life), add it at the very end of your reply as "
        "[remember: one short fact in third person]. One tag per fact, only new facts, never things already in the notes. "
        "If something in the notes stops being true, add [forget: a few words from that line] and, if needed, [remember: the new fact]. "
        "Don't save small talk or moods that pass. The tags are hidden from them.\n\n"
        "What you remember about them:\n" + (notes or "(nothing yet: you're just getting to know them)") + "\n\n"
        + brief_for_ai(db, user, link.tz_minutes)
    )
    try:
        reply = await ai.chat(turns, system, 2400 if file_text else 1500)
    except ai.AIError as exc:
        log.warning("logbook: telegram companion failed: %s", str(exc)[:200])
        return "My brain blanked for a sec 😅 try me again in a minute?"
    save_turn(db, user.id, "user", said)
    reply = apply_memory(db, user.id, reply, local_today(link.tz_minutes))
    save_turn(db, user.id, "assistant", reply)
    return esc(reply.replace("**", ""))


async def handle_command(db: Session, user: User, link: TelegramLink, cmd: str, arg: str) -> None:
    chat = link.chat_id
    tz = link.tz_minutes
    today = local_today(tz)
    uid = user.id

    if cmd in ("help", "start"):
        await send(chat, HELP)
        return

    if cmd == "memory":
        notes = get_notes(db, uid)
        if not notes:
            await send(chat, "I don't have anything saved about you yet 🌱 Tell me about your schools, scores and plans and I'll keep them.")
        else:
            lines = [l.split(": ", 1)[-1] for l in notes.splitlines() if l.strip()]
            await send_long(chat, "🧠 <b>What I remember about you</b>\n\n" + "\n".join("• " + esc(l) for l in lines) +
                            "\n\n/forget <i>words</i> drops a line, /forget everything wipes it all.")
        return

    if cmd == "forget":
        if not arg:
            await send(chat, "What should I forget? /forget <i>a few words from the line</i>, or /forget everything.")
            return
        if arg.lower() in ("everything", "all"):
            set_notes(db, uid, "")
            db.query(CompanionMessage).filter(CompanionMessage.user_id == uid).delete()
            db.commit()
            await send(chat, "Done, clean slate 🫧 I forgot everything and our chat history.")
            return
        lines = [l for l in get_notes(db, uid).splitlines() if l.strip()]
        keep = [l for l in lines if arg.lower() not in l.lower()]
        set_notes(db, uid, "\n".join(keep))
        gone = len(lines) - len(keep)
        await send(chat, f"Forgot {gone} thing{'s' if gone != 1 else ''} ✨" if gone else "I couldn't find that in what I remember 🤔 /memory shows it all.")
        return

    if cmd == "log":
        if not arg:
            await send(chat, "Log time like this:\n<code>/log 45m #coding fixed the login bug</code>\n<code>/log 1h30 #study SAT reading</code>")
            return
        title, area, minutes = parse_log(arg)
        if not title:
            title = AREAS.get(area or "", "Logged time")
        ensure_day(db, uid, today)
        act = Activity(user_id=uid, day_date=today, title=title[:300], activity_category=area, duration_minutes=minutes)
        db.add(act)
        db.commit()
        total = sum(a.duration_minutes or 0 for a in db.query(Activity).filter(Activity.user_id == uid, Activity.day_date == today))
        what = (minutes_label(minutes) + " of " if minutes else "") + (AREAS[area].lower() if area else "time")
        text = f"🕒 Logged {what}: <b>{esc(title)}</b>" + (f"\nToday so far: {minutes_label(total)}. {cheer(total)}" if total else "")
        buttons = None
        if not area:
            text += "\nWhich area?"
            names = list(AREAS)
            buttons = [[{"text": AREAS[a], "callback_data": f"a:{act.id}:{a}"} for a in names[i:i + 3]] for i in range(0, 9, 3)]
        await send(chat, text, buttons)
        return

    if cmd in ("task", "t", "tomorrow"):
        day = today + timedelta(days=1) if cmd == "tomorrow" else today
        if not arg:
            await send(chat, f"What's the task? <code>/{cmd} call the bank</code>")
            return
        title, area, _ = parse_log(arg) if "#" in arg else (arg, None, None)
        ensure_day(db, uid, day)
        db.add(Task(user_id=uid, day_date=day, title=(title or arg)[:300], activity_category=area))
        db.commit()
        tasks = tasks_for(db, uid, day)
        label = "tomorrow" if cmd == "tomorrow" else "today"
        await send(chat, f"Added to {label}: <b>{esc(title or arg)}</b>\n" + tasks_text(tasks, label), task_buttons(tasks))
        return

    if cmd == "tasks":
        tasks = tasks_for(db, uid, today)
        await send(chat, tasks_text(tasks, "today"), task_buttons(tasks))
        return

    if cmd in ("j", "journal", "entry"):
        if not arg:
            await send(chat, "Write it after the command: <code>/j today felt long but good</code>")
            return
        entry = Entry(user_id=uid, occurred_at=datetime.now(timezone.utc), mood="okay", text=arg[:20000])
        db.add(entry)
        db.commit()
        await send(chat, "📝 Got it, saved to your journal. How did today feel?",
                   [[{"text": label, "callback_data": f"m:{entry.id}:{mood}"} for mood, label in MOODS.items()]])
        return

    if cmd == "dump":
        if not arg:
            await send(chat, "Empty your head: <code>/dump call grandma, buy notebook</code>")
            return
        rec = ensure_day(db, uid, today)
        rec.brain_dump = (rec.brain_dump.rstrip() + "\n" if rec.brain_dump.strip() else "") + arg.strip()
        db.commit()
        await send(chat, "🧠 Dropped it in today's brain dump. Lighter now?")
        return

    if cmd == "pray":
        if not arg:
            await send(chat, "What are you praying for? <code>/pray wisdom for the applications</code>")
            return
        now = datetime.now(timezone.utc)
        db.add(Prayer(user_id=uid, text=arg.strip()[:4000], created_at=now, updated_at=now))
        db.commit()
        await send(chat, "🙏 Added to your prayers. Praying with you.")
        return

    if cmd == "prayers":
        rows = db.query(Prayer).filter(Prayer.user_id == uid, Prayer.answered_at.is_(None)).order_by(Prayer.created_at.desc()).limit(15).all()
        if not rows:
            await send(chat, "No open prayers. Add one: /pray ...")
            return
        await send(chat, "<b>Praying for</b>\n" + "\n".join("• " + esc(p.text) for p in rows) + "\n\nTap one when it's answered.",
                   [[{"text": "🙌 " + p.text[:50], "callback_data": f"p:{p.id}"}] for p in rows])
        return

    if cmd in ("due", "deadline"):
        at, arg2 = parse_time(arg)
        when, rest = parse_date(arg2, today)
        if not arg or when is None or not rest:
            await send(chat, "Add a date and what's due:\n<code>/due nov 1 23:59 Early decision</code>\n<code>/due in 3 weeks SAT</code>\n"
                             "<code>/due 2026-12-05 9am SAT #exam</code>\n<code>/due mar 4 2009 Sara #birthday</code>")
            return
        if re.search(r"#?birthday\b|\bbday\b", rest, re.I):
            kind = "birthday"
        elif re.search(r"#exam\b|\bexam\b|\btest\b|\bsat\b", rest, re.I):
            kind = "exam"
        elif re.search(r"#event\b", rest, re.I):
            kind = "event"
        else:
            kind = "deadline"
        title = re.sub(r"#(exam|event|deadline)\b|(?:'s)?\s*#?(?:birthday|bday)\b", "", rest, flags=re.I).strip(" ,-") or rest
        row = Deadline(user_id=uid, title=title[:300], due_date=when, due_time=at, kind=kind)
        db.add(row)
        db.commit()
        row.next_date = occurrence(row, today)
        await send(chat, f"{KIND_ICON[kind]} <b>{esc(title_for(row))}</b> on {row.next_date.strftime('%a, %b')} {row.next_date.day}"
                         f"{' at ' + at if at else ''}: {countdown((row.next_date - today).days)}.")
        return

    if cmd == "deadlines":
        rows = upcoming(db, uid, today, 365)
        if not rows:
            await send(chat, "Nothing coming up. Add one: /due nov 1 Early decision")
            return
        await send(chat, "<b>Coming up</b>\n" + "\n".join(
            f"{KIND_ICON.get(d.kind, '🚩')} {esc(title_for(d))} · {d.next_date.strftime('%b')} {d.next_date.day}{' ' + d.due_time if d.due_time else ''} · {countdown((d.next_date - today).days)}" for d in rows[:15]))
        return

    if cmd == "habits":
        rows = db.query(Habit).filter(Habit.user_id == uid, Habit.archived.is_(False)).order_by(Habit.sort, Habit.id).all()
        if not rows:
            await send(chat, "No habits yet. Add them on the Today page in Logbook.")
            return
        done = {h for (h,) in db.query(HabitLog.habit_id).filter(HabitLog.user_id == uid, HabitLog.date == today)}
        await send(chat, f"<b>Habits</b> · {len(done & {h.id for h in rows})} of {len(rows)} today",
                   [[{"text": ("✅ " if h.id in done else "⬜ ") + h.name[:56], "callback_data": f"h:{h.id}"}] for h in rows])
        return

    if cmd == "today":
        await send(chat, summary(db, user, tz))
        return

    if cmd == "verse":
        ref = verses.ref_for(today)
        words = await verses.words(ref)
        await send(chat, f"📖 <i>{esc(words)}</i>\n— {esc(ref)}" if words else f"📖 Today's verse is {esc(ref)}.")
        return

    if cmd == "morning":
        a = arg.strip().lower()
        if a in ("off", "stop", "no"):
            link.morning_hour = None
            msg = "Morning message is off. /morning 7 turns it back on."
        else:
            m = re.match(r"^(\d{1,2})(?::00)?\s*(am|pm)?$", a)
            hour = 7 if not a else (int(m.group(1)) % 12 + (12 if m.group(2) == "pm" else 0) if m and m.group(2) else int(m.group(1)) if m else None)
            if hour is None or not 0 <= hour <= 23:
                await send(chat, "Pick an hour: <code>/morning 7</code>, <code>/morning 6am</code> or <code>/morning off</code>")
                return
            link.morning_hour = hour
            msg = f"☀️ I'll message you every morning at {hour}:00."
        db.commit()
        refresh_links(db)
        await send(chat, msg)
        return

    if cmd == "sleep":
        night = today + timedelta(days=1) if local_now(tz).hour >= 12 else today
        rec = db.get(SleepLog, (uid, night)) or SleepLog(user_id=uid, date=night)
        rec.bed_at, rec.wake_at = datetime.now(timezone.utc), None
        db.merge(rec)
        db.commit()
        await send(chat, "🌙 Goodnight! You did enough today. Send /wake when you're up and I'll tell you how long you slept.")
        return

    if cmd == "wake":
        rows = (db.query(SleepLog).filter(SleepLog.user_id == uid, SleepLog.date >= today - timedelta(days=1),
                                          SleepLog.bed_at.is_not(None), SleepLog.wake_at.is_(None)).all())
        if not rows:
            await send(chat, "Good morning ☀️ I don't have a bedtime for last night. /sleep at night next time, or log it in the app.")
            return
        rec = rows[-1]
        rec.wake_at = datetime.now(timezone.utc)
        bed = rec.bed_at if rec.bed_at.tzinfo else rec.bed_at.replace(tzinfo=timezone.utc)
        mins = round((rec.wake_at - bed).total_seconds() / 60)
        if not 0 < mins <= 20 * 60:
            rec.wake_at = None
            db.commit()
            await send(chat, "That doesn't look like one night's sleep. Log it in the app instead?")
            return
        db.commit()
        await send(chat, f"☀️ Morning! You slept <b>{minutes_label(mins)}</b>." + (" That's a good night 😌" if mins >= 420 else " A bit short, so go easy on yourself today. Early night later?"))
        return

    if cmd in ("grateful", "thanks", "gratitude"):
        if not arg:
            await send(chat, "What are you grateful for? <code>/grateful family dinner</code>")
            return
        rec = ensure_day(db, uid, today)
        rec.reflection_gratitude = (rec.reflection_gratitude.rstrip() + "\n" if rec.reflection_gratitude.strip() else "") + arg.strip()
        db.commit()
        total = sum(len([l for l in (r.reflection_gratitude or "").splitlines() if l.strip()])
                    for r in db.query(DayRecord).filter(DayRecord.user_id == uid))
        await send(chat, f"✨ In the jar. That's {total} firefl{'y' if total == 1 else 'ies'}.")
        return

    if cmd in ("nudges", "reminders"):
        off = arg.strip().lower() in ("off", "stop", "no")
        set_setting(db, f"tg_nudges_off:{uid}", "1" if off else "")
        db.commit()
        nudge_plans.pop(uid, None)
        await send(chat, "Reminders are off. /nudges on brings them back." if off else
                   "🔔 Reminders on: deadlines the day before and 2 hours before, and your streak at 21:00 if you haven't written.")
        return

    if cmd == "backup":
        a = arg.strip().lower()
        if a in ("off", "stop", "no"):
            set_setting(db, f"tg_backup_off:{uid}", "1")
            db.commit()
            await send(chat, "Weekly backups are off. /backup on turns them back on; /backup sends one now.")
            return
        if a in ("on", "yes"):
            set_setting(db, f"tg_backup_off:{uid}", "")
            db.commit()
            await send(chat, "🗄 I'll send a backup every Sunday morning.")
            return
        if not await send_backup(db, user, chat, tz):
            await send(chat, "Couldn't send the backup just now. Try again in a minute.")
        return

    if cmd == "unlink":
        db.delete(link)
        db.commit()
        refresh_links(db)
        await send(chat, "Disconnected. Your Logbook is untouched. Connect again from Logbook's settings.")
        return

    await send(chat, "Hmm, I don't know that one 😅 /help has everything I can do.")


async def handle_callback(db: Session, user: User, link: TelegramLink, cb: dict) -> None:
    chat = link.chat_id
    data = cb.get("data", "")
    msg = cb.get("message") or {}
    message_id = msg.get("message_id")
    uid = user.id
    today = local_today(link.tz_minutes)
    note = ""
    kind, _, rest = data.partition(":")

    if kind == "t":
        task = db.get(Task, int(rest)) if rest.isdigit() else None
        if task and task.user_id == uid:
            task.completed = not task.completed
            db.commit()
            tasks = tasks_for(db, uid, task.day_date)
            label = "today" if task.day_date == today else "tomorrow" if task.day_date == today + timedelta(days=1) else task.day_date.isoformat()
            if message_id:
                await edit(chat, message_id, tasks_text(tasks, label), task_buttons(tasks))
            note = "All done for today. Good work 🎉" if tasks and all(t.completed for t in tasks) and label == "today" else ("Done ✅" if task.completed else "Unchecked")

    elif kind == "m":
        entry_id, _, mood = rest.partition(":")
        entry = db.get(Entry, (uid, entry_id))
        if entry and mood in MOODS:
            entry.mood = mood
            db.commit()
            if message_id:
                await edit(chat, message_id, f"📝 Saved to your journal · {MOODS[mood]}")
            note = {"good": "Glad it was a good one.", "okay": "Noted.", "rough": "Sorry it was rough. I'm here."}[mood]

    elif kind == "a":
        act_id, _, area = rest.partition(":")
        act = db.get(Activity, int(act_id)) if act_id.isdigit() else None
        if act and act.user_id == uid and area in AREAS:
            act.activity_category = area
            db.commit()
            if message_id:
                await edit(chat, message_id, f"🕒 Logged to {AREAS[area].lower()}: <b>{esc(act.title)}</b>")
            note = AREAS[area]

    elif kind == "h":
        habit = db.get(Habit, int(rest)) if rest.isdigit() else None
        if habit and habit.user_id == uid:
            logged = db.get(HabitLog, (habit.id, today))
            if logged:
                db.delete(logged)
            else:
                db.add(HabitLog(user_id=uid, habit_id=habit.id, date=today))
            db.commit()
            rows = db.query(Habit).filter(Habit.user_id == uid, Habit.archived.is_(False)).order_by(Habit.sort, Habit.id).all()
            done = {h for (h,) in db.query(HabitLog.habit_id).filter(HabitLog.user_id == uid, HabitLog.date == today)}
            if message_id:
                await edit(chat, message_id, f"<b>Habits</b> · {len(done & {h.id for h in rows})} of {len(rows)} today",
                           [[{"text": ("✅ " if h.id in done else "⬜ ") + h.name[:56], "callback_data": f"h:{h.id}"}] for h in rows])
            note = "Nice 🌱" if not logged else "Unchecked"

    elif kind == "p":
        prayer = db.get(Prayer, int(rest)) if rest.isdigit() else None
        if prayer and prayer.user_id == uid and prayer.answered_at is None:
            prayer.answered_at = datetime.now(timezone.utc)
            db.commit()
            note = "Answered 🙌 Worth remembering."
            await send(chat, f"🙌 Marked answered: <b>{esc(prayer.text)}</b>\nAdd how it was answered in Logbook, under Prayers.")

    try:
        await call(state["token"], "answerCallbackQuery", {"callback_query_id": cb.get("id"), "text": note[:180]})
    except TelegramError:
        pass


async def process(update: dict) -> None:
    """One update from Telegram, start to finish. Never raises."""
    try:
        with SessionLocal() as db:
            if not load_token(db):
                return
            if links_cache is None:
                refresh_links(db)
            if "callback_query" in update:
                cb = update["callback_query"]
                chat_id = ((cb.get("message") or {}).get("chat") or {}).get("id")
                link = db.query(TelegramLink).filter(TelegramLink.chat_id == chat_id).first() if chat_id else None
                user = db.get(User, link.user_id) if link else None
                if user:
                    await handle_callback(db, user, link, cb)
                return

            msg = update.get("message") or {}
            chat = msg.get("chat") or {}
            media = msg.get("photo") or msg.get("document")
            if chat.get("type") != "private" or ("text" not in msg and not media):
                return
            chat_id = chat["id"]
            text = (msg.get("text") or "").strip()
            link = db.query(TelegramLink).filter(TelegramLink.chat_id == chat_id).first()

            cmd, arg = "", text
            m = re.match(r"^/(\w+)(?:@\w+)?\s*(.*)$", text, re.S)
            if m:
                cmd, arg = m.group(1).lower(), m.group(2).strip()

            if cmd == "start" and arg:
                found = codes.pop(arg, None)
                if not found or found[2] < time.time():
                    await send(chat_id, "That link has expired. Make a new one in Logbook: Settings → Telegram.")
                    return
                user_id, tz, _ = found
                db.query(TelegramLink).filter((TelegramLink.chat_id == chat_id) | (TelegramLink.user_id == user_id)).delete()
                db.add(TelegramLink(user_id=user_id, chat_id=chat_id, tz_minutes=tz, morning_hour=7))
                db.commit()
                refresh_links(db)
                user = db.get(User, user_id)
                await send(chat_id, f"Connected to <b>{esc(user.username if user else 'your')}</b>'s Logbook ✨ I'm {esc(pet_name(db, user_id))}.\n\n" + HELP +
                           "\n\nI'll send a short morning message at 7:00. /morning off to stop it.")
                return

            if link is None:
                await send(chat_id, "This is a private Logbook bot. To connect, open Logbook → Settings → Telegram → Connect.")
                return
            user = db.get(User, link.user_id)
            if user is None:
                return
            if msg.get("document"):
                await handle_document(db, user, link, msg)
                return
            if media:
                await handle_media(db, user, link, msg)
                return
            if cmd:
                await handle_command(db, user, link, cmd, arg)
            else:
                try:
                    await call(state["token"], "sendChatAction", {"chat_id": chat_id, "action": "typing"})
                except TelegramError:
                    pass
                await send_long(chat_id, await companion_reply(db, user, link, text))
    except Exception:  # one bad update must never take the server down
        log.exception("logbook: telegram update failed")


def spawn(coro) -> None:
    task = asyncio.create_task(coro)
    _tasks.add(task)
    task.add_done_callback(_tasks.discard)


# ---------- running in the background ----------

async def morning_tick(now: datetime | None = None) -> int:
    """Sends the morning message to everyone whose hour it is. Returns how many went out."""
    if not links_cache:
        return 0
    now = now or datetime.now(timezone.utc)
    sent = 0
    for uid, link in list(links_cache.items()):
        hour = link["hour"]
        if hour is None:
            continue
        local = now + timedelta(minutes=link["tz"])
        if local.hour != hour or link["last"] == local.date():
            continue
        link["last"] = local.date()
        with SessionLocal() as db:
            row = db.get(TelegramLink, uid)
            user = db.get(User, uid)
            if row is None or user is None:
                continue
            row.last_morning = local.date()
            db.commit()
            await send(row.chat_id, await morning_message(db, user, row.tz_minutes))
            if local.weekday() == 6 and weekly_backup_on(db, uid):     # Sundays: a copy of everything
                await send_backup(db, user, row.chat_id, row.tz_minutes)
            if local.day == 1:                                          # the 1st: the letter about last month
                await send_letter(db, user, row.chat_id, row.tz_minutes, local.date())
            sent += 1
    return sent


# ---------- reminders through the day ----------

nudge_plans: dict[int, dict] = {}     # user id -> {"built": time, "items": [(when utc, key, kind, text)]}
nudges_sent: set[str] = set()


def nudges_on(db: Session, uid: int) -> bool:
    return get_setting(db, f"tg_nudges_off:{uid}") != "1"


def build_nudges(db: Session, uid: int, tz: int, now: datetime) -> list:
    """A deadline the day before and two hours before, and the streak at 21:00."""
    items = []
    local = now + timedelta(minutes=tz)
    for d in upcoming(db, uid, local.date(), 2):
        hh, mm = (int(x) for x in (d.due_time or "23:59").split(":"))
        due_utc = datetime.combine(d.next_date, datetime.min.time(), timezone.utc).replace(hour=hh, minute=mm) - timedelta(minutes=tz)
        name = esc(title_for(d))
        items.append((due_utc - timedelta(hours=24), f"{uid}:d{d.id}:{d.next_date}:24", "due",
                      f"{KIND_ICON.get(d.kind, '🚩')} <b>{name}</b> is tomorrow" + (f" at {d.due_time}." if d.due_time else ".")))
        items.append((due_utc - timedelta(hours=2), f"{uid}:d{d.id}:{d.next_date}:2", "due",
                      f"⏳ <b>{name}</b>: 2 hours left." + (" You've got this." if d.kind != "birthday" else "")))
    evening = datetime.combine(local.date(), datetime.min.time(), timezone.utc).replace(hour=21) - timedelta(minutes=tz)
    items.append((evening, f"{uid}:s:{local.date()}", "streak", ""))
    return items


async def nudge_tick(now: datetime | None = None) -> int:
    if not links_cache:
        return 0
    now = now or datetime.now(timezone.utc)
    sent = 0
    for uid, link in list(links_cache.items()):
        plan = nudge_plans.get(uid)
        if plan is None or (now - plan["built"]).total_seconds() > 3600 or plan["day"] != (now + timedelta(minutes=link["tz"])).date():
            with SessionLocal() as db:
                if not nudges_on(db, uid):
                    nudge_plans[uid] = {"built": now, "day": (now + timedelta(minutes=link["tz"])).date(), "items": []}
                    continue
                nudge_plans[uid] = plan = {"built": now, "day": (now + timedelta(minutes=link["tz"])).date(),
                                           "items": build_nudges(db, uid, link["tz"], now)}
        for when, key, kind, text in plan["items"]:
            if key in nudges_sent or not (timedelta(0) <= now - when < timedelta(minutes=10)):
                continue
            nudges_sent.add(key)
            if kind == "streak":
                with SessionLocal() as db:
                    today = local_today(link["tz"])
                    wrote = db.query(Entry).filter(Entry.user_id == uid).all()
                    if any(((e.occurred_at if e.occurred_at.tzinfo else e.occurred_at.replace(tzinfo=timezone.utc))
                            + timedelta(minutes=link["tz"])).date() == today for e in wrote):
                        continue
                    n = streak(db, uid, today, link["tz"])
                    text = (f"🌙 Your <b>{n}-day streak</b> is waiting. One line is enough: /j ..." if n
                            else "🌙 A line before bed? /j how today went")
            await send(link["chat"], text)
            sent += 1
    if len(nudges_sent) > 5000:
        nudges_sent.clear()
    return sent


async def poll_forever() -> None:
    offset = 0
    while state["token"] and state["mode"] == "polling":
        try:
            updates = await call(state["token"], "getUpdates", {"offset": offset, "timeout": 25,
                                                                "allowed_updates": ["message", "callback_query"]}, timeout=40)
        except TelegramError as exc:
            log.warning("logbook: telegram polling: %s", exc)
            await asyncio.sleep(15)
            continue
        for u in updates or []:
            offset = max(offset, u.get("update_id", 0) + 1)
            await process(u)


async def run() -> None:
    """Started with the server: loads the bot, polls if there's no webhook,
    and sends morning messages. Quiet (and cheap) when no bot is set up."""
    try:
        with SessionLocal() as db:
            load_token(db)
            state["mode"] = get_setting(db, "telegram_mode")
            if state["token"]:
                refresh_links(db)
    except Exception:
        log.exception("logbook: telegram start failed")
        return
    poller = None
    while True:
        if state["token"] and state["mode"] == "polling" and (poller is None or poller.done()):
            poller = asyncio.create_task(poll_forever())
        try:
            await morning_tick()
        except Exception:
            log.exception("logbook: morning message failed")
        try:
            await nudge_tick()
        except Exception:
            log.exception("logbook: reminders failed")
        await asyncio.sleep(60)
